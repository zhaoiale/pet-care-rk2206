#!/usr/bin/env python3
"""生成智宠管家演示 PPT (4 页)"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ── 通用颜色 ──
BG_DARK   = RGBColor(0x1A, 0x1A, 0x2E)
BG_CARD   = RGBColor(0x28, 0x2A, 0x3E)
ACCENT    = RGBColor(0x4C, 0xAF, 0x50)
ACCENT2   = RGBColor(0x21, 0x96, 0xF3)
ACCENT3   = RGBColor(0xFF, 0x98, 0x00)
RED       = RGBColor(0xF4, 0x43, 0x36)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
GRAY      = RGBColor(0xAA, 0xAA, 0xAA)
LGRAY     = RGBColor(0xCC, 0xCC, 0xCC)

FONT_TITLE = 'Microsoft YaHei'
FONT_BODY  = 'Microsoft YaHei'
FONT_MONO  = 'Consolas'


def set_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, font_size=18,
                color=WHITE, bold=False, font_name=FONT_BODY, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return tf


def add_card(slide, left, top, width, height):
    """添加一个圆角矩形卡片"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG_CARD
    shape.line.fill.background()
    return shape


def add_accent_line(slide, left, top, width, color=ACCENT):
    """添加一条强调线"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(0.04)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


# ═══════════════════════════════════════
# 第 1 页：背景 & 痛点
# ═══════════════════════════════════════
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # blank
set_bg(slide1, BG_DARK)

# 装饰线
add_accent_line(slide1, 0, 0, 13.333, ACCENT)

# 主标题
add_textbox(slide1, 1.0, 1.0, 11.3, 1.0, '宠物独处焦虑 — 智宠管家',
            font_size=40, bold=True, font_name=FONT_TITLE)

# 副标题
add_textbox(slide1, 1.0, 2.0, 11.3, 0.6,
            '基于 OpenHarmony RK2206 的多模态感知与主动关怀系统',
            font_size=20, color=GRAY)

add_accent_line(slide1, 1.0, 2.8, 2.0, ACCENT)

# 痛点区域：左侧
add_textbox(slide1, 1.0, 3.3, 5.5, 0.5, '▎背景', font_size=22, bold=True, color=ACCENT)

pain_points = [
    '中国宠物数量超 1.2 亿，主人日均离家 8 小时以上',
    '独处焦虑 → 破坏行为、过度吠叫、食欲异常',
    '市面产品多为被动监测，缺乏"感知→安抚"闭环',
]
for i, txt in enumerate(pain_points):
    add_textbox(slide1, 1.0, 3.9 + i * 0.55, 5.5, 0.5, f'•  {txt}',
                font_size=16, color=LGRAY)

# 右侧：目标
add_textbox(slide1, 7.5, 3.3, 5.0, 0.5, '▎核心目标', font_size=22, bold=True, color=ACCENT2)

goals = [
    '多传感器融合感知宠物状态',
    '焦虑评估 → 主动分级安抚',
    'App 远程投喂 + 数据看板',
    'GPS 电子围栏防走失',
]
for i, txt in enumerate(goals):
    add_textbox(slide1, 7.5, 3.9 + i * 0.55, 5.0, 0.5, f'✦  {txt}',
                font_size=16, color=LGRAY)

# 底部 Slogan
add_accent_line(slide1, 1.0, 6.5, 11.3, ACCENT)
add_textbox(slide1, 1.0, 6.7, 11.3, 0.5,
            '感知 → 评估 → 安抚 → 闭环',
            font_size=24, bold=True, color=ACCENT, alignment=PP_ALIGN.CENTER)


# ═══════════════════════════════════════
# 第 2 页：系统架构
# ═══════════════════════════════════════
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide2, BG_DARK)
add_accent_line(slide2, 0, 0, 13.333, ACCENT)

add_textbox(slide2, 1.0, 0.6, 11.3, 0.8, '系统架构',
            font_size=36, bold=True, font_name=FONT_TITLE)
add_accent_line(slide2, 1.0, 1.4, 2.0, ACCENT)

# App 层
app_card = add_card(slide2, 1.0, 1.8, 11.3, 1.2)
add_textbox(slide2, 1.3, 1.95, 10.7, 0.5, 'HarmonyOS App (ArkTS)',
            font_size=22, bold=True, color=ACCENT2)
add_textbox(slide2, 1.3, 2.45, 10.7, 0.4,
            'Dashboard | 投喂管理 | GPS 轨迹 | 安抚控制',
            font_size=16, color=LGRAY)

# MQTT 层
add_textbox(slide2, 1.0, 3.15, 11.3, 0.5, 'MQTT  (HiveMQ 公共 Broker)',
            font_size=16, color=GRAY, alignment=PP_ALIGN.CENTER)

# 设备层
dev_card = add_card(slide2, 1.0, 3.7, 11.3, 2.2)
add_textbox(slide2, 1.3, 3.85, 10.7, 0.5, 'RK2206  LiteOS-M 主控',
            font_size=22, bold=True, color=ACCENT)

# 传感器网格
sensors = [
    ('MPU6050', '活动量'),
    ('MAX30102', '心率/血氧'),
    ('HX711', '食盆称重'),
    ('ATGM336H', 'GPS 定位'),
    ('PIR + 麦克', '活动/叫声'),
    ('BH1750', '光照强度'),
    ('SHT30', '温湿度'),
    ('SU-03T', '语音交互'),
]
for i, (name, desc) in enumerate(sensors):
    col = i % 4
    row = i // 4
    x = 1.4 + col * 2.75
    y = 4.45 + row * 0.7
    add_card(slide2, x, y, 2.45, 0.55)
    add_textbox(slide2, x + 0.1, y + 0.05, 2.25, 0.25,
                name, font_size=13, bold=True, color=ACCENT3)
    add_textbox(slide2, x + 0.1, y + 0.3, 2.25, 0.25,
                desc, font_size=11, color=GRAY)

# 执行层
add_card(slide2, 1.0, 6.1, 11.3, 0.7)
add_textbox(slide2, 1.3, 6.2, 10.7, 0.5,
            '继电器投喂  |  LED 安抚灯  |  蜂鸣器音频安抚  |  LCD 信息显示',
            font_size=16, color=LGRAY, alignment=PP_ALIGN.CENTER)


# ═══════════════════════════════════════
# 第 3 页：核心创新点
# ═══════════════════════════════════════
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide3, BG_DARK)
add_accent_line(slide3, 0, 0, 13.333, ACCENT)

add_textbox(slide3, 1.0, 0.6, 11.3, 0.8, '核心创新点',
            font_size=36, bold=True, font_name=FONT_TITLE)
add_accent_line(slide3, 1.0, 1.4, 2.0, ACCENT)

innovations = [
    {
        'num': '①',
        'title': '多模态焦虑评估模型',
        'desc': 'MPU6050 活动量 + MAX30102 心率 + HRV 压力 + 叫声检测 + 环境温湿度',
        'highlight': '6 维特征 → 加权融合 → 焦虑等级 (0-100)',
        'color': ACCENT2,
    },
    {
        'num': '②',
        'title': 'HX711 闭环克级投喂',
        'desc': 'App 设定目标克数 → 继电器启动电机 → HX711 实时称重 → 达到减量目标自动停止',
        'highlight': '精度 ±10g | 60s 超时保护 | 无传感器时估算降级',
        'color': ACCENT,
    },
    {
        'num': '③',
        'title': '分级主动安抚策略',
        'desc': '轻度焦虑 (25-44) → 暖光 LED  |  中度 (45-69) → 蜂鸣器音乐旋律  |  重度 (70+) → 主人录音安抚',
        'highlight': '感知 → 评估 → 安抚闭环，焦虑实时回落可见',
        'color': ACCENT3,
    },
    {
        'num': '④',
        'title': '电子围栏 + GPS 离线轨迹',
        'desc': 'ATGM336H 定位 → Flash 三扇区 256 点离线存储 → MQTT 10 点/包上传 → App 回放',
        'highlight': '任意时刻设家为围栏中心，越界立即推送告警',
        'color': RED,
    },
]

for i, item in enumerate(innovations):
    y = 1.8 + i * 1.3
    card = add_card(slide3, 1.0, y, 11.3, 1.1)
    # 编号
    add_textbox(slide3, 1.3, y + 0.15, 0.5, 0.5, item['num'],
                font_size=26, bold=True, color=item['color'])
    # 标题 + 描述
    add_textbox(slide3, 2.0, y + 0.12, 9.5, 0.4, item['title'],
                font_size=22, bold=True)
    add_textbox(slide3, 2.0, y + 0.52, 9.5, 0.35, item['desc'],
                font_size=14, color=LGRAY)
    # 高亮
    add_textbox(slide3, 2.0, y + 0.82, 9.5, 0.3, item['highlight'],
                font_size=13, bold=True, color=item['color'])


# ═══════════════════════════════════════
# 第 4 页：过渡 → 演示
# ═══════════════════════════════════════
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide4, BG_DARK)
add_accent_line(slide4, 0, 0, 13.333, ACCENT)

# 居中大标题
add_textbox(slide4, 0, 2.0, 13.333, 1.5, '演 示',
            font_size=72, bold=True, font_name=FONT_TITLE, alignment=PP_ALIGN.CENTER)

add_accent_line(slide4, 4.5, 3.6, 4.3, ACCENT)

add_textbox(slide4, 0, 4.0, 13.333, 0.8,
            'RK2206 开发板实物  +  HarmonyOS App 完整数据流',
            font_size=24, color=GRAY, alignment=PP_ALIGN.CENTER)

# 底部清单
items = [
    'LCD 卡片式实时仪表盘',
    'App 7 天历史数据看板',
    '远程投喂 → HX711 闭环控制',
    '焦虑评估 → 分级安抚联动',
]
for i, txt in enumerate(items):
    x = 1.5 + i * 2.8
    add_textbox(slide4, x, 5.3, 2.6, 0.5, txt,
                font_size=14, bold=True, color=LGRAY, alignment=PP_ALIGN.CENTER)
    if i < len(items) - 1:
        add_textbox(slide4, x + 2.6, 5.35, 0.2, 0.4, '|',
                    font_size=14, color=GRAY, alignment=PP_ALIGN.CENTER)

# ── 保存 ──
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), '智宠管家_演示PPT.pptx')
prs.save(output_path)
print(f'PPT saved to: {output_path}')
